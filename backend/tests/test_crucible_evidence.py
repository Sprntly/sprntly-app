"""`app.crucible.evidence` — the substrate a later selection stage will read:
when a table is FROM, its SHAPE without its rows, and the two deterministic
retrieval marks. No selection happens here, and nothing here is acted on.

Pure: no DB, no LLM. The one real `.xlsx` round-trip below is a real
`openpyxl` write/read against `tmp_path`, exercising the actual reader
(`recon.tables_from_workbook`) rather than a hand-built `Table`.
"""
from __future__ import annotations

from datetime import datetime, timezone

from app.crucible import evidence
from app.crucible.recon import make_table, tables_from_workbook


def _workbook(path, sheets: dict[str, tuple[list[str], list[list]]]) -> None:
    """An `.xlsx` on disk — same shape as `test_crucible_chat_uploads._workbook`,
    kept local so this file exercises the real reader without importing
    another test module's fixture helper."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(str(path))


_NOW = datetime(2026, 9, 8, tzinfo=timezone.utc)


# ─── 1. Schema view: one entry per sheet, correct counts, never a row ───────


def test_schema_view_has_one_entry_per_sheet_with_correct_counts_and_kinds(
    tmp_path,
):
    path = tmp_path / "workbook.xlsx"
    _workbook(path, {
        "contracts": (
            ["account", "acv_usd", "closed_on"],
            [["Account A", 22500, "2025-01-15"],
             ["Account B", 372810, "2025-02-01"]],
        ),
        "notes": (
            ["field", "value"],
            [["source", "CRM export"], ["owner", "Revenue Ops"]],
        ),
    })
    tables = tables_from_workbook(path)
    snapshot = evidence.schema_view(tables)

    assert len(snapshot.tables) == 2
    by_name = {t.name.split(":")[-1]: t for t in snapshot.tables}
    assert by_name["contracts"].row_count == 2
    assert by_name["notes"].row_count == 2

    contracts_cols = {c.name: c for c in by_name["contracts"].columns}
    assert contracts_cols["acv_usd"].kind == "number"
    assert contracts_cols["account"].kind == "text"
    assert contracts_cols["closed_on"].kind == "date"


def test_schema_view_never_carries_a_row():
    table = make_table(
        "sheet:t", [{"a": "1", "b": "2"}, {"a": "3", "b": "4"}],
        columns=["a", "b"],
    )
    snapshot = evidence.schema_view([table])

    assert not hasattr(snapshot, "rows")
    for t in snapshot.tables:
        assert not hasattr(t, "rows")
        for c in t.columns:
            assert not hasattr(c, "rows")
    import json
    assert '"rows"' not in json.dumps(snapshot.to_json())


def test_schema_view_exposes_a_workspace_today_without_filtering_on_it():
    table = make_table("sheet:t", [{"a": "1"}], columns=["a"])
    snapshot = evidence.schema_view([table], today=_NOW)
    assert snapshot.today == _NOW
    # Every table still appears — no staleness gate anywhere in this module.
    assert len(snapshot.tables) == 1


# ─── 2. as_of: all three rules, and which one fired ─────────────────────────


def test_as_of_rule_1_is_the_max_plausible_value_in_a_date_column():
    table = make_table(
        "deals:sheet1",
        [{"account": "A", "closed_on": "2024-06-01", "amt": 100},
         {"account": "B", "closed_on": "2025-01-15", "amt": 200}],
        columns=["account", "closed_on", "amt"],
    )
    as_of, rule = evidence.derive_as_of(table, now=_NOW)
    assert rule == evidence.AS_OF_RULE_TABLE_DATE_COLUMN
    assert (as_of.year, as_of.month, as_of.day) == (2025, 1, 15)


def test_as_of_rule_2_resolves_a_month_year_column_header_to_that_month():
    """The worked case from the spec: a column header that names the month
    and year of the figures beside it, with no date column anywhere."""
    table = make_table(
        "pricing:sheet1",
        [{"competitor": "Acme", "Est. price (Nov 2025)": 199}],
        columns=["competitor", "Est. price (Nov 2025)"],
    )
    as_of, rule = evidence.derive_as_of(table, now=_NOW)
    assert rule == evidence.AS_OF_RULE_HEADER_OR_NAME
    assert (as_of.year, as_of.month) == (2025, 11)


def test_as_of_rule_2_resolves_an_iso_date_in_the_sheet_name():
    table = make_table(
        "competitor_pricing_2025-11-03:sheet1",
        [{"item": "plan a", "price": 10}], columns=["item", "price"],
    )
    as_of, rule = evidence.derive_as_of(table, now=_NOW)
    assert rule == evidence.AS_OF_RULE_HEADER_OR_NAME
    assert (as_of.year, as_of.month, as_of.day) == (2025, 11, 3)


def test_as_of_rule_3_falls_back_to_the_upload_timestamp():
    table = make_table(
        "misc:sheet1", [{"a": "x", "b": "y"}], columns=["a", "b"],
    )
    as_of, rule = evidence.derive_as_of(table, now=_NOW)
    assert rule == evidence.AS_OF_RULE_UPLOAD_TIMESTAMP
    assert as_of == _NOW


def test_as_of_rule_3_prefers_an_explicit_upload_timestamp_over_now():
    table = make_table(
        "misc:sheet1", [{"a": "x", "b": "y"}], columns=["a", "b"],
    )
    stamp = datetime(2026, 8, 1, tzinfo=timezone.utc)
    as_of, rule = evidence.derive_as_of(table, now=_NOW, upload_timestamp=stamp)
    assert rule == evidence.AS_OF_RULE_UPLOAD_TIMESTAMP
    assert as_of == stamp


def test_with_as_of_hydrates_tables_without_touching_rows_or_columns():
    table = make_table(
        "deals:sheet1", [{"account": "A", "closed_on": "2025-01-15"}],
        columns=["account", "closed_on"],
    )
    hydrated = evidence.with_as_of([table], now=_NOW)
    assert len(hydrated) == 1
    out = hydrated[0]
    assert out.as_of is not None
    assert out.as_of_rule == evidence.AS_OF_RULE_TABLE_DATE_COLUMN
    assert out.rows == table.rows
    assert out.columns == table.columns
    assert out.source_type == table.source_type
    # The original is untouched — `Table` is frozen and `replace` copies.
    assert table.as_of is None
    assert table.as_of_rule == ""


def test_derive_as_of_never_raises_on_a_table_with_nothing_to_date():
    table = make_table("nothing:sheet1", [], columns=[])
    as_of, rule = evidence.derive_as_of(table, now=_NOW)
    assert rule == evidence.AS_OF_RULE_UPLOAD_TIMESTAMP
    assert as_of == _NOW


# ─── 3. passthrough_sheet ────────────────────────────────────────────────────


def test_passthrough_sheet_marks_an_eight_row_field_value_sheet():
    rows = [{"field": f"f{i}", "value": f"v{i}"} for i in range(8)]
    table = make_table("methodology:sheet1", rows, columns=["field", "value"])
    assert evidence.is_passthrough_sheet(table) is True


def test_passthrough_sheet_does_not_mark_a_350_row_table():
    rows = [
        {"account": f"Account {i}", "mrr": i * 10, "region": "us",
         "tier": "gold"}
        for i in range(350)
    ]
    table = make_table(
        "accounts:sheet1", rows, columns=["account", "mrr", "region", "tier"],
    )
    assert evidence.is_passthrough_sheet(table) is False


def test_passthrough_sheet_row_bound_is_named_and_matches_the_worked_case():
    """The constant IS the eight-row worked case's headroom, not an
    accident — an 8-row sheet must clear it and a 21-row one must not."""
    at_bound = make_table(
        "sheet:t",
        [{"a": str(i)} for i in range(evidence.PASSTHROUGH_SHEET_MAX_ROWS)],
        columns=["a"],
    )
    over_bound = make_table(
        "sheet:t",
        [{"a": str(i), "b": str(i * 2), "c": str(i * 3)}
         for i in range(evidence.PASSTHROUGH_SHEET_MAX_ROWS + 1)],
        columns=["a", "b", "c"],
    )
    assert evidence.is_passthrough_sheet(at_bound) is True
    assert evidence.is_passthrough_sheet(over_bound) is False


def test_passthrough_sheet_marks_a_long_field_value_shaped_table():
    """A field/value sheet running longer than the row bound is still marked
    — the shape matters, not just the row count."""
    rows = [{"field": f"field_{i}", "value": f"value_{i}"} for i in range(40)]
    table = make_table("glossary:sheet1", rows, columns=["field", "value"])
    assert evidence.is_passthrough_sheet(table) is True


# ─── 4. passthrough_column ───────────────────────────────────────────────────


_LOSS_NOTES = (
    "Customer said the renewal price was well above budget and they found a "
    "cheaper competitor offering similar core features this cycle.",
    "They cited cost pressure from a recent budget freeze and could not "
    "justify the increase given current usage levels this quarter.",
    "Went with a rival vendor that had a more mature integration story for "
    "their existing internal analytics stack and reporting tools.",
    "Finance flagged the renewal as too expensive relative to actual usage "
    "and asked the team to revisit pricing before renewing at all.",
)


def test_passthrough_column_marks_freetext_beside_a_coded_column():
    rows = [
        {"loss_reason": "price", "loss_notes": _LOSS_NOTES[0]},
        {"loss_reason": "price", "loss_notes": _LOSS_NOTES[1]},
        {"loss_reason": "competitor", "loss_notes": _LOSS_NOTES[2]},
        {"loss_reason": "price", "loss_notes": _LOSS_NOTES[3]},
    ]
    table = make_table(
        "deals:losses", rows, columns=["loss_reason", "loss_notes"],
    )
    assert evidence.passthrough_columns(table) == ("loss_notes",)


def test_passthrough_column_does_not_mark_freetext_with_no_coded_sibling():
    rows = [{"notes": text} for text in _LOSS_NOTES]
    table = make_table("deals:notes_only", rows, columns=["notes"])
    assert evidence.passthrough_columns(table) == ()


def test_passthrough_column_marks_nothing_when_no_column_is_freetext():
    rows = [
        {"loss_reason": "price", "region": "us"},
        {"loss_reason": "competitor", "region": "eu"},
        {"loss_reason": "price", "region": "us"},
    ]
    table = make_table("deals:coded_only", rows, columns=["loss_reason", "region"])
    assert evidence.passthrough_columns(table) == ()


# ─── 5. Marks are recorded, never acted on ───────────────────────────────────


def test_schema_view_carries_the_marks_without_altering_row_count_or_columns():
    rows = [{"field": f"f{i}", "value": f"v{i}"} for i in range(8)]
    table = make_table("methodology:sheet1", rows, columns=["field", "value"])
    snapshot = evidence.schema_view([table])
    t = snapshot.tables[0]
    assert t.passthrough_sheet is True
    assert t.row_count == 8
    assert [c.name for c in t.columns] == ["field", "value"]
