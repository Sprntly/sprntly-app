"""A file attached in chat, read by a Goal Analysis plan — and written nowhere.

THE WHOLE POINT IS THE THING THAT DOES NOT HAPPEN. The same spreadsheet can
reach a run through Settings → Connectors → Uploads, which INGESTS it: rows in
`kg_signal`, permanent, tenant-wide, answering every later question and not
removable from the product. For a company's standing sources that is correct.
For a file someone is analysing once it is a one-way contamination of the one
asset the product cannot rebuild — which is exactly why the run-scoped path
exists, and why `test_nothing_reaches_the_knowledge_graph` is the load-bearing
test in this file rather than a formality at the bottom of it.

The rest establishes that the bytes actually arrive: through the composer
(which refused spreadsheets), past the per-message cap (which refused twelve
files), through a key that is checked against the caller's own workspace, into
the SAME `observe()` call the connected corpus goes into — because every
interesting check here is cross-table and observing the two sets separately
would silently discard every finding that spans them.

Pure and offline: workbooks are built in `tmp_path`, storage is monkeypatched.
"""
from __future__ import annotations

import pytest

from app import attachments_storage
from app.crucible import plan as plan_mod
from app.crucible import recon
from app.crucible.planner import _evidence_sentence


# ─── Workbooks, built rather than committed ─────────────────────────────────


def _workbook(path, sheets: dict[str, tuple[list[str], list[list]]]) -> bytes:
    """An .xlsx on disk and its bytes, so one fixture serves both the
    read-from-disk and the read-from-storage paths."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(str(path))
    return path.read_bytes()


@pytest.fixture
def contracts_xlsx(tmp_path):
    """`_tabular_recon_fixtures.contracts()`, SERIALISED TO A REAL WORKBOOK.

    Reusing the canonical fixture rather than inventing a second table with
    the same intent: that shape is the one `test_crucible_recon.py` pins the
    reconciliation check against, so an upload built from it is guaranteed to
    produce a real observation — and if the check's thresholds ever move, this
    file moves with it instead of quietly becoming a test that asserts a table
    was read and nothing more.
    """
    from tests import _tabular_recon_fixtures as fx

    table = fx.contracts()
    data = _workbook(tmp_path / "08_sales_data.xlsx", {
        "contracts": (
            list(table.columns),
            [[r.get(c) for c in table.columns] for r in table.rows],
        ),
    })
    return tmp_path / "08_sales_data.xlsx", data, len(table.rows)


# ─── 1. The composer's own guards, which refused the file at the door ───────


def test_a_workbook_is_a_supported_attachment_extension():
    """`stage_attachment` raises on an unsupported extension and the upload
    route 422s before that — so without these entries the bytes never reached
    storage at all, and every path below is unreachable."""
    assert attachments_storage.is_supported_ext("xlsx")
    assert attachments_storage.is_supported_ext("xls")
    assert attachments_storage.media_type_for_key("k/a.xlsx") == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_a_message_may_carry_more_than_eight_files():
    """A twelve-file evidence pack is one ordinary send. Both send routes are
    asserted because they are one composer to the person using them, and a cap
    raised on only one of them is a cap that still refuses the send."""
    from app.routes.ask import AskIn
    from app.routes.conversations import TurnIn

    for model in (AskIn, TurnIn):
        cap = model.model_fields["attachments"].metadata
        limits = [getattr(m, "max_length", None) for m in cap]
        assert 16 in limits, f"{model.__name__} still caps attachments lower"


# ─── 2. A key is an authorisation claim, and is checked as one ──────────────


def test_a_key_from_another_workspace_is_not_owned():
    assert attachments_storage.owns_key(
        workspace_id="ws-1", key="chat-attachments/ws-1/a.xlsx")
    assert not attachments_storage.owns_key(
        workspace_id="ws-1", key="chat-attachments/ws-2/a.xlsx")


def test_a_traversal_segment_does_not_become_a_read():
    """`..` can satisfy the prefix and still address a path outside it, which
    is why the prefix check alone is not the whole check."""
    assert not attachments_storage.owns_key(
        workspace_id="ws-1", key="chat-attachments/ws-1/../ws-2/a.xlsx")
    with pytest.raises(ValueError):
        attachments_storage.read_attachment(
            workspace_id="ws-1", key="chat-attachments/ws-1/../ws-2/a.xlsx")


def test_reading_another_workspaces_key_raises_rather_than_returning_bytes():
    with pytest.raises(ValueError):
        attachments_storage.read_attachment(
            workspace_id="ws-1", key="chat-attachments/ws-2/a.xlsx")


def test_a_staged_file_reads_back_byte_identical(tmp_path, monkeypatch):
    monkeypatch.setattr(attachments_storage, "_bucket_name", lambda: None)
    monkeypatch.setattr(attachments_storage.settings, "storage_dir",
                        str(tmp_path), raising=False)
    key = "chat-attachments/ws-1/abc.xlsx"
    attachments_storage._stage_filesystem_sync(key, b"PK\x03\x04payload")
    assert attachments_storage.read_attachment(
        workspace_id="ws-1", key=key) == b"PK\x03\x04payload"


# ─── 3. The name is the reader's; the parser is the key's ───────────────────


def test_the_filename_survives_so_the_plan_can_name_the_upload():
    assert recon.upload_filename("08_sales_data.xlsx",
                                 "chat-attachments/w/uuid.xlsx") == \
        "08_sales_data.xlsx"


def test_a_name_carrying_directories_is_reduced_to_a_leaf():
    """The name is client-supplied and reaches a filesystem, so it may only
    ever be a leaf."""
    for hostile in ("../../etc/passwd.xlsx", "/etc/passwd.xlsx",
                    r"..\..\windows\sys.xlsx"):
        out = recon.upload_filename(hostile, "chat-attachments/w/uuid.xlsx")
        assert "/" not in out and "\\" not in out, out
        assert not out.startswith("."), out


def test_the_extension_comes_from_the_key_not_from_the_name():
    """The key's extension is the one the WRITE side validated. Taking it from
    the name would let a caller choose which reader runs over the bytes."""
    assert recon.upload_filename("book.csv", "chat-attachments/w/u.xlsx") == \
        "book.xlsx"


def test_a_name_that_reduces_to_nothing_falls_back_to_the_key():
    out = recon.upload_filename("...", "chat-attachments/w/uuid.xlsx")
    assert out == "uuid.xlsx"


# ─── 4. Bytes in memory produce the same tables as bytes on disk ────────────


def test_an_upload_read_from_storage_matches_the_same_file_read_from_disk(
    contracts_xlsx,
):
    """The property that makes the whole feature trustworthy: what the plan
    saw cannot depend on how the bytes arrived."""
    path, data, _n = contracts_xlsx
    from_disk = recon.tables_from_dir(path.parent,
                                      source_type=recon.UPLOAD_SOURCE_TYPE)
    from_storage = recon.tables_from_uploads([("08_sales_data.xlsx", data)])

    assert [t.name for t in from_disk] == [t.name for t in from_storage]
    assert [t.columns for t in from_disk] == [t.columns for t in from_storage]
    assert [list(t.rows) for t in from_disk] == [list(t.rows) for t in from_storage]
    assert ({o.kind for o in recon.observe(from_disk).observations}
            == {o.kind for o in recon.observe(from_storage).observations})
    # And it is not a vacuous match of two empty lists.
    assert from_storage
    assert "value_columns_disagree" in {
        o.kind for o in recon.observe(from_storage).observations}


def test_two_attachments_with_the_same_filename_both_survive(contracts_xlsx):
    """A second write under the same name would replace the first, and the
    plan would report one table where the reader sent two."""
    _, data, n_rows = contracts_xlsx
    tables = recon.tables_from_uploads(
        [("book.xlsx", data), ("book.xlsx", data)])
    assert len(tables) == 2
    assert len({t.name for t in tables}) == 2


def test_a_file_that_is_not_tabular_costs_its_own_table_and_nothing_else(
    contracts_xlsx,
):
    """The pack is mostly PDFs. They carry prose, not columns, and a plan that
    fell over on one — or claimed to have read it — would be worse than one
    that reads the spreadsheets and says so."""
    _, data, n_rows = contracts_xlsx
    tables = recon.tables_from_uploads([
        ("00_readme.pdf", b"%PDF-1.4 not a workbook"),
        ("08_sales_data.xlsx", data),
        ("notes.docx", b"PK\x03\x04 not a workbook either"),
    ])
    assert [t.name for t in tables] == ["08_sales_data:contracts"]


def test_an_empty_attachment_is_skipped():
    assert recon.tables_from_uploads([("empty.xlsx", b"")]) == []


# ─── 5. Uploads join the corpus in ONE observe() call ───────────────────────


def _recon_env(monkeypatch, *, signals, uploads_bytes, signal_error=None):
    """Wire `_recon_report` to fixed signals and fixed attachment bytes."""
    import app.routes.crucible as routes

    def _page(_client, _company_id, page):
        return list(signals) if page == 0 else []

    def _boom(*_a, **_k):
        raise RuntimeError("statement timeout")

    monkeypatch.setattr(routes, "_signal_page", _boom if signal_error else _page)
    monkeypatch.setattr("app.db.client.require_client", lambda: object())
    monkeypatch.setattr(
        attachments_storage, "read_attachment",
        lambda *, workspace_id, key: uploads_bytes[key],
    )
    return routes


def test_the_upload_and_the_corpus_are_observed_together(
    monkeypatch, contracts_xlsx,
):
    """ONE CALL, NOT TWO. Every check worth having here is cross-table, so
    observing the upload and the graph separately would hand the reader the
    findings inside each and silently drop every finding that spans them."""
    _, data, n_rows = contracts_xlsx
    key = "chat-attachments/ws-1/u.xlsx"
    routes = _recon_env(monkeypatch, signals=[], uploads_bytes={key: data})

    seen = {}
    real_observe = recon.observe

    def _spy(tables, **kw):
        seen["tables"] = list(tables)
        return real_observe(tables, **kw)

    monkeypatch.setattr(recon, "observe", _spy)
    report = routes._recon_report(
        "co-1", uploads=((key, "08_sales_data.xlsx"),), workspace_id="ws-1")

    assert "08_sales_data:contracts" in {t.name for t in seen["tables"]}
    assert "value_columns_disagree" in {o.kind for o in report.observations}


def test_a_run_with_no_attachments_reads_exactly_what_it_read_before(
    monkeypatch,
):
    """The corpus query is untouched by this feature: uploads are additive."""
    routes = _recon_env(monkeypatch, signals=[], uploads_bytes={})
    assert routes._recon_report("co-1").to_json() == recon.ReconReport().to_json()


def test_a_failed_corpus_read_does_not_take_the_uploads_with_it(
    monkeypatch, contracts_xlsx,
):
    """A statement timeout on the signal query used to mean an empty report.
    A reader who attached twelve files would be shown a plan that read
    nothing, while everything they handed over sat in storage, readable."""
    _, data, n_rows = contracts_xlsx
    key = "chat-attachments/ws-1/u.xlsx"
    routes = _recon_env(monkeypatch, signals=[], uploads_bytes={key: data},
                        signal_error=True)
    report = routes._recon_report(
        "co-1", uploads=((key, "08_sales_data.xlsx"),), workspace_id="ws-1")
    assert "value_columns_disagree" in {o.kind for o in report.observations}


def test_an_unreadable_attachment_costs_its_table_and_not_the_plan(
    monkeypatch, contracts_xlsx,
):
    _, data, n_rows = contracts_xlsx
    good, gone = "chat-attachments/ws-1/a.xlsx", "chat-attachments/ws-1/b.xlsx"
    import app.routes.crucible as routes

    monkeypatch.setattr(routes, "_signal_page", lambda *a, **k: [])
    monkeypatch.setattr("app.db.client.require_client", lambda: object())

    def _read(*, workspace_id, key):
        if key == gone:
            raise RuntimeError("object missing")
        return data

    monkeypatch.setattr(attachments_storage, "read_attachment", _read)
    report = routes._recon_report(
        "co-1", uploads=((good, "kept.xlsx"), (gone, "swept.xlsx")),
        workspace_id="ws-1")
    assert [s.name for s in report.sources] == ["kept:contracts"]


def test_nothing_reaches_the_knowledge_graph(monkeypatch, contracts_xlsx):
    """THE CONSTRAINT THIS FEATURE EXISTS FOR.

    The tenant this is tested with has real Fireflies data in its graph. The
    connector-upload path would write this pack into `kg_signal` permanently
    and there is no undo in the product. So: no write of any kind, to any
    table, on the reconnaissance path — asserted by making every mutating
    verb on the client raise, rather than by inspecting what was called.
    """
    _, data, n_rows = contracts_xlsx
    key = "chat-attachments/ws-1/u.xlsx"

    class _Forbidden:
        def __init__(self, name):
            self._name = name

        def __getattr__(self, verb):
            if verb in ("insert", "upsert", "update", "delete", "rpc"):
                raise AssertionError(
                    f"the run-scoped upload path wrote to {self._name!r} "
                    f"via .{verb}() — nothing may reach the knowledge graph")
            return lambda *a, **k: self

        def execute(self):
            class _R:
                data: list = []
                count = 0
            return _R()

    class _Client:
        def table(self, name):
            return _Forbidden(name)

        def __getattr__(self, item):
            raise AssertionError(f"unexpected client access: {item}")

    import app.routes.crucible as routes

    monkeypatch.setattr("app.db.client.require_client", _Client)
    monkeypatch.setattr(
        attachments_storage, "read_attachment",
        lambda *, workspace_id, key: data)

    report = routes._recon_report(
        "co-1", uploads=((key, "08_sales_data.xlsx"),), workspace_id="ws-1")
    # And it did the work — a vacuously empty report would pass the above.
    assert "value_columns_disagree" in {o.kind for o in report.observations}


# ─── 6. The plan says the files are being read, and says they are its own ───


def _upload_report(data, names=("08_sales_data.xlsx",)):
    return recon.observe(recon.tables_from_uploads([(n, data) for n in names]))


def test_an_upload_is_listed_as_a_source_named_by_its_filename(contracts_xlsx):
    _, data, n_rows = contracts_xlsx
    ups = plan_mod.uploads_from_report(_upload_report(data))
    assert [u.name for u in ups] == ["08_sales_data"]
    assert ups[0].tables == 1
    assert ups[0].records == n_rows


def test_several_sheets_of_one_workbook_are_one_upload(tmp_path):
    """A workbook is one thing to the person who attached it; counting its
    sheets as several is how a plan claims eleven sources over four files."""
    data = _workbook(tmp_path / "analytics.xlsx", {
        "funnel": (["stage", "n"], [["a", 3], ["b", 2]]),
        "cohorts": (["month", "n"], [["2026-01", 9], ["2026-02", 8]]),
    })
    ups = plan_mod.uploads_from_report(_upload_report(data, ("analytics.xlsx",)))
    assert len(ups) == 1
    assert ups[0].tables == 2
    assert ups[0].records == 4


def test_a_file_that_produced_no_table_is_not_claimed_as_read():
    """The plan lists what it READ, not what it was handed. Echoing the twelve
    filenames would tell a reader it read seven PDFs it could not open."""
    report = recon.observe(recon.tables_from_uploads(
        [("00_readme.pdf", b"%PDF-1.4")]))
    assert plan_mod.uploads_from_report(report) == ()


def test_a_connected_source_is_never_listed_as_an_upload():
    connected = recon.make_table(
        "gong_calls", [{"a": 1}, {"a": 2}], source_type="customer_voice")
    assert plan_mod.uploads_from_report(recon.observe([connected])) == ()


def test_the_opening_step_names_the_attached_files_and_scopes_them(
    contracts_xlsx,
):
    """A reader is being asked to APPROVE a method. If its first sentence puts
    their upload among their connected sources, they approve on a false
    premise — and the difference matters, because the same question asked
    tomorrow without the file gets a different answer."""
    _, data, n_rows = contracts_xlsx
    sentence = _evidence_sentence(_upload_report(data), ())
    assert "sales data" in sentence
    assert "attached to this message" in sentence
    assert "not added to your knowledge graph" in sentence


def test_a_run_over_connected_sources_alone_says_what_it_always_said():
    """The negative twin. Without it, the assertions above are satisfied by a
    sentence that says 'attached' over every corpus in the product."""
    connected = recon.make_table(
        "gong_calls", [{"a": 1}, {"a": 2}], source_type="customer_voice")
    sentence = _evidence_sentence(recon.observe([connected]), ())
    assert sentence.endswith("traced back to something you connected.")
    assert "attached" not in sentence
    assert "knowledge graph" not in sentence


def test_the_plan_carries_the_uploads_into_its_stored_json(contracts_xlsx):
    """The gate renders from this blob, and the approve path carries it
    forward — an upload absent here is an upload the reader never sees."""
    _, data, n_rows = contracts_xlsx
    built = plan_mod.RunPlan(
        goal_text="g", definition_text="d", currency="accounts",
        uploads=plan_mod.uploads_from_report(_upload_report(data)),
    )
    assert built.to_json()["uploads"] == [
        {"name": "08_sales_data", "tables": 1, "records": n_rows}]


def test_a_plan_built_before_this_existed_still_renders():
    assert plan_mod.RunPlan(
        goal_text="g", definition_text="d", currency="accounts",
    ).to_json()["uploads"] == []


# ─── 7. A file that was NOT read is named, never silently absent ────────────
#
# ONE ATTACHMENT PER RUN, THROUGHOUT THIS SECTION. A run carrying twelve files
# cannot attribute a failure to any one of them: the symptom that produced
# these tests was three files missing from a twelve-file pack, and the only
# reason it was ever diagnosable is that the same files attached ALONE
# succeeded. A test that asserts an aggregate count would have passed against
# the broken engine as readily as the fixed one.


def _corpus_signals(n_tables: int) -> list[dict]:
    """Signals shaped into exactly `n_tables` structured tables.

    `tables_from_signals` buckets by (source_type, kind) and keeps a bucket
    with two or more rows, so this is the smallest corpus that occupies a
    given number of table slots. The real tenant this reproduces carries
    11,402 signals and yields enough buckets to fill the budget on its own.
    """
    out: list[dict] = []
    for i in range(n_tables):
        for row in range(2):
            out.append({
                "id": f"s{i}-{row}", "source_type": "analytics",
                "kind": f"metric_{i:03d}",
                "properties": {"value": float(row), "delta": float(i)},
            })
    return out


def test_one_unreadable_attachment_is_named_in_the_plan_rather_than_omitted(
    monkeypatch,
):
    """THE INVARIANT, AT ITS SMALLEST. One file in, nothing readable out — and
    the reader is told which file and why, instead of being handed a plan that
    mentions no attachment at all and reads as though none was sent."""
    key = "chat-attachments/ws-1/a.pdf"
    routes = _recon_env(monkeypatch, signals=[],
                        uploads_bytes={key: b"%PDF-1.4 prose, not columns"})
    report = routes._recon_report(
        "co-1", uploads=((key, "09_crm_win_loss.pdf"),), workspace_id="ws-1")

    unread = plan_mod.unread_uploads_from_report(report)
    assert [u.name for u in unread] == ["09_crm_win_loss"]
    assert "spreadsheets" in unread[0].reason
    # And it is not claimed as read in the same breath.
    assert plan_mod.uploads_from_report(report) == ()


def test_one_attachment_swept_from_storage_is_named_rather_than_dropped(
    monkeypatch,
):
    """The fetch that fails is the one failure the reader can actually act on
    — re-attach the file — and it was the most silent of the lot."""
    import app.routes.crucible as routes

    monkeypatch.setattr(routes, "_signal_page", lambda *a, **k: [])
    monkeypatch.setattr("app.db.client.require_client", lambda: object())

    def _read(*, workspace_id, key):
        raise RuntimeError("object missing")

    monkeypatch.setattr(attachments_storage, "read_attachment", _read)
    report = routes._recon_report(
        "co-1", uploads=(("chat-attachments/ws-1/b.xlsx", "10_nps_responses.xlsx"),),
        workspace_id="ws-1")

    unread = plan_mod.unread_uploads_from_report(report)
    assert [u.name for u in unread] == ["10_nps_responses"]
    assert "storage" in unread[0].reason


def test_one_attachment_is_still_read_when_the_corpus_fills_the_budget(
    monkeypatch, contracts_xlsx,
):
    """THE ROOT CAUSE, AS A TEST.

    `observe` reads at most `MAX_TABLES` tables and the reconnaissance pass
    used to hand it the knowledge-graph tables first, so on a tenant whose
    graph fills the budget the reader's own attachment was truncated off the
    end — silently, with the plan then listing whatever survived as though it
    were everything sent. Measured on staging: this exact file attached ALONE
    was read correctly and lost a sheet when it followed others.

    The corpus here is sized to fill the budget on its own, which is what
    makes the assertion new-path-only: before the fix there is no room left
    and the upload appears nowhere.
    """
    _, data, n_rows = contracts_xlsx
    key = "chat-attachments/ws-1/u.xlsx"
    routes = _recon_env(monkeypatch,
                        signals=_corpus_signals(recon.MAX_TABLES),
                        uploads_bytes={key: data})
    report = routes._recon_report(
        "co-1", uploads=((key, "09_crm_win_loss.xlsx"),), workspace_id="ws-1")

    assert [u.name for u in plan_mod.uploads_from_report(report)] == [
        "09_crm_win_loss"]
    # Read in full, not partly: a file that keeps two of its three sheets is
    # the same defect one sheet smaller.
    assert plan_mod.uploads_from_report(report)[0].records == n_rows
    assert plan_mod.unread_uploads_from_report(report) == ()
    # And the corpus was TRIMMED, not thrown away: the two populations still
    # go into one `observe` call, which is what the cross-table checks need.
    assert any(s.source_type == "analytics" for s in report.sources)


def test_a_sheet_that_does_not_fit_even_first_is_named_rather_than_lost(
    monkeypatch, tmp_path,
):
    """The bound still exists, and now it speaks. One workbook carrying more
    sheets than the whole pass reads keeps as many as fit and SAYS how many it
    did not — which is the difference between a stated limit and a file that
    quietly came back smaller than it went in."""
    sheets = {
        f"sheet{i:02d}": (["stage", "n"], [["a", i], ["b", i + 1]])
        for i in range(recon.MAX_TABLES + 3)
    }
    data = _workbook(tmp_path / "03_product_analytics.xlsx", sheets)
    key = "chat-attachments/ws-1/big.xlsx"
    routes = _recon_env(monkeypatch, signals=[], uploads_bytes={key: data})
    report = routes._recon_report(
        "co-1", uploads=((key, "03_product_analytics.xlsx"),),
        workspace_id="ws-1")

    unread = plan_mod.unread_uploads_from_report(report)
    assert [u.name for u in unread] == ["03_product_analytics"]
    assert "3 of its sheets" in unread[0].reason
    # It is listed on BOTH lists on purpose: partly read is neither "read" nor
    # "not read", and a reader deciding whether to approve needs both facts.
    assert [u.name for u in plan_mod.uploads_from_report(report)] == [
        "03_product_analytics"]


def test_the_plan_carries_what_it_could_not_read_into_its_stored_json():
    """The gate renders from this blob. A disclosure absent here is a
    disclosure the reader never sees."""
    built = plan_mod.RunPlan(
        goal_text="g", definition_text="d", currency="accounts",
        unread_uploads=(plan_mod.UnreadUpload(
            name="11_third_party_feedback", reason="it could not be opened"),),
    )
    assert built.to_json()["unread_uploads"] == [
        {"name": "11_third_party_feedback", "reason": "it could not be opened"}]


def test_a_plan_built_before_this_existed_reads_back_as_nothing_unread():
    assert plan_mod.RunPlan(
        goal_text="g", definition_text="d", currency="accounts",
    ).to_json()["unread_uploads"] == []
